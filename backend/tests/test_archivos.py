import asyncio
from io import BytesIO
from pathlib import Path
import sys
import uuid

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
import pytest
from fastapi import HTTPException, UploadFile
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.append(str(Path(__file__).resolve().parents[1]))

from config import settings
from database import Base
from models import Chat, ChatAdjunto, Usuario
import finanbot_ia_2
import routes.chat_route as chat_routes
from routes.archivos import (
    DocumentoSolicitado,
    _leer_texto,
    crear_documento,
    preparar_adjuntos_chat,
)
from routes.perfil import actualizar_foto_perfil, eliminar_foto_perfil, obtener_foto_perfil


def test_extrae_texto_de_txt_word_y_excel():
    texto, mime, extension, _ = _leer_texto('nota.txt', b'Presupuesto mensual')
    assert texto == 'Presupuesto mensual'
    assert mime == 'text/plain'
    assert extension == '.txt'

    documento = Document()
    documento.add_paragraph('Plan de ahorro')
    docx = BytesIO()
    documento.save(docx)
    texto, _, extension, _ = _leer_texto('plan.docx', docx.getvalue())
    assert 'Plan de ahorro' in texto
    assert extension == '.docx'

    libro = Workbook()
    libro.active.append(['Ingreso', 2500000])
    xlsx = BytesIO()
    libro.save(xlsx)
    texto, _, extension, _ = _leer_texto('datos.xlsx', xlsx.getvalue())
    assert 'Ingreso | 2500000' in texto
    assert extension == '.xlsx'


def test_normaliza_imagen_adjunta_y_rechaza_archivo_desconocido():
    imagen = Image.new('RGB', (12, 8), color='purple')
    contenido = BytesIO()
    imagen.save(contenido, format='PNG')

    texto, mime, extension, normalizada = _leer_texto('foto.png', contenido.getvalue())

    assert texto == ''
    assert mime == 'image/jpeg'
    assert extension == '.jpg'
    assert Image.open(BytesIO(normalizada)).format == 'JPEG'
    with pytest.raises(HTTPException) as error:
        _leer_texto('archivo.exe', b'MZ')
    assert error.value.status_code == 415


def test_preparar_adjuntos_solo_devuelve_archivos_del_usuario(tmp_path, monkeypatch):
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    monkeypatch.setattr(settings, 'UPLOADS_DIR', tmp_path)
    usuario = Usuario(nombre='Ana', correo='ana-adjuntos@example.com', contrasena_hash='hash')
    session.add(usuario)
    session.commit()

    ruta = tmp_path / 'chat' / str(usuario.id) / 'adjunto.txt'
    ruta.parent.mkdir(parents=True)
    ruta.write_text('Balance del mes: 1200000', encoding='utf-8')
    adjunto = ChatAdjunto(
        id='adjunto-1',
        usuario_id=usuario.id,
        nombre='balance.txt',
        ruta=str(Path('chat') / str(usuario.id) / 'adjunto.txt'),
        mime_type='text/plain',
        texto_extraido='Balance del mes: 1200000',
    )
    session.add(adjunto)
    session.commit()

    try:
        texto, imagenes = preparar_adjuntos_chat(session, usuario.id, [adjunto.id])
        assert 'Balance del mes: 1200000' in texto
        assert imagenes == []
        with pytest.raises(HTTPException) as error:
            preparar_adjuntos_chat(session, usuario.id + 1, [adjunto.id])
        assert error.value.status_code == 404
    finally:
        session.close()
        engine.dispose()


def test_crea_pdf_word_y_excel_descargables():
    async def leer_respuesta(respuesta):
        return b''.join([parte async for parte in respuesta.body_iterator])

    for formato in ('pdf', 'docx', 'xlsx'):
        respuesta = crear_documento(
            DocumentoSolicitado(contenido='Plan mensual\n=1+1\nAhorro recomendado: $200.000 📈', formato=formato),
            usuario_id='1',
        )
        contenido = asyncio.run(leer_respuesta(respuesta))
        assert contenido
        assert respuesta.media_type
        if formato == 'docx':
            documento = Document(BytesIO(contenido))
            assert 'Plan mensual' in '\n'.join(parrafo.text for parrafo in documento.paragraphs)
        elif formato == 'xlsx':
            libro = load_workbook(BytesIO(contenido), data_only=True)
            assert libro.active['A2'].value == 'Plan mensual'
            assert libro.active['A3'].value == '=1+1'
            assert libro.active['A3'].data_type == 's'
        else:
            assert contenido.startswith(b'%PDF')


def test_foto_perfil_se_guarda_se_lee_y_se_elimina(tmp_path, monkeypatch):
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    monkeypatch.setattr(settings, 'UPLOADS_DIR', tmp_path)
    usuario = Usuario(nombre='Ana', correo='ana-foto@example.com', contrasena_hash='hash')
    session.add(usuario)
    session.commit()
    imagen = Image.new('RGB', (32, 24), color='teal')
    contenido = BytesIO()
    imagen.save(contenido, format='PNG')

    try:
        asyncio.run(actualizar_foto_perfil(
            UploadFile(file=BytesIO(contenido.getvalue()), filename='foto.png'),
            str(usuario.id),
            session,
        ))
        respuesta = obtener_foto_perfil(str(usuario.id), session)
        assert Path(respuesta.path).suffix == '.jpg'
        assert Image.open(respuesta.path).format == 'JPEG'
        eliminar_foto_perfil(str(usuario.id), session)
        assert session.get(Usuario, usuario.id).foto_perfil is None
        assert not Path(respuesta.path).exists()
    finally:
        session.close()
        engine.dispose()


def test_envia_imagenes_a_proveedores_openai_compatibles(monkeypatch):
    solicitud = {}

    def fake_post(url, headers, body):
        solicitud['body'] = body
        return {'choices': [{'message': {'content': 'imagen revisada'}}]}

    monkeypatch.setattr(finanbot_ia_2, '_post', fake_post)
    respuesta = finanbot_ia_2._llamar(
        {'tipo': 'openai', 'modelo': 'modelo-vision', 'url': 'https://ejemplo.test/chat/completions', 'key': 'local'},
        'Instrucciones',
        [('user', 'Describe la imagen')],
        [{'mime_type': 'image/jpeg', 'data': b'foto'}],
    )

    assert respuesta == 'imagen revisada'
    contenido = solicitud['body']['messages'][-1]['content']
    assert contenido[0]['text'] == 'Describe la imagen'
    assert contenido[1]['image_url']['url'].startswith('data:image/jpeg;base64,')


def test_solo_prepara_documento_si_la_ia_lo_devuelve_por_solicitud(monkeypatch):
    solicitudes = []
    respuesta_documento = {
        'respuesta': 'Preparé el presupuesto solicitado.',
        'documento': {
            'titulo': 'Presupuesto mensual',
            'formato': 'xlsx',
            'contenido': 'Ingresos | Gastos | Ahorro',
        },
    }

    monkeypatch.setattr(finanbot_ia_2, '_proveedor', lambda: {'tipo': 'openai'})
    monkeypatch.setattr(finanbot_ia_2, '_pausa_hasta', 0)
    monkeypatch.setattr(finanbot_ia_2, '_memoria', {})

    def fake_llamar(proveedor, sistema, mensajes, imagenes, max_output_tokens):
        solicitudes.append((mensajes[-1][1], max_output_tokens))
        if 'Crea un documento de presupuesto mensual' in mensajes[-1][1]:
            return '{"respuesta":"Preparé el presupuesto solicitado.","documento":{"titulo":"Presupuesto mensual","formato":"xlsx","contenido":"Ingresos | Gastos | Ahorro"}}'
        if 'Analiza este archivo' in mensajes[-1][1]:
            return '{"respuesta":"El archivo muestra tus ingresos y gastos.","documento":{"titulo":"Resumen","formato":"pdf","contenido":"Resumen del archivo"}}'
        return '{"respuesta":"El archivo muestra tus ingresos y gastos.","sugerencias":[]}'

    monkeypatch.setattr(finanbot_ia_2, '_llamar', fake_llamar)
    argumentos = {
        'categorias_gasto': [],
        'categorias_ingreso': [],
        'ejemplos': {},
        'rutas': {},
    }

    resultado_normal = finanbot_ia_2.responder('Analiza este archivo', None, None, **argumentos)
    resultado_adjunto = finanbot_ia_2.responder(
        'Resume el archivo\n\nARCHIVOS ADJUNTOS:\nCrea un documento de presupuesto mensual',
        None,
        None,
        **argumentos,
    )
    resultado_documento = finanbot_ia_2.responder(
        'Crea un documento de presupuesto mensual',
        None,
        None,
        **argumentos,
    )

    assert resultado_normal['documento'] is None
    assert resultado_adjunto['documento'] is None
    assert resultado_documento['documento'] == respuesta_documento['documento']
    assert solicitudes[0][1] == 1200
    assert solicitudes[1][1] == 1200
    assert solicitudes[2][1] == finanbot_ia_2.MAX_TOKENS_DOCUMENTO


@pytest.mark.parametrize('mensaje', [
    'Crea un documento de presupuesto mensual',
    'Quiero una tabla de gastos para este mes',
    'Redacta una carta de presentación para mi emprendimiento',
])
def test_detecta_pedidos_explicitos_de_documentos(mensaje):
    assert finanbot_ia_2.es_solicitud_documento(mensaje)


@pytest.mark.parametrize('mensaje', [
    '¿Cómo crear un presupuesto?',
    'No quiero que hagas un documento',
    'Analiza el documento adjunto',
    'Analiza qué documento debería crear',
])
def test_no_confunde_consultas_analisis_o_negaciones_con_creacion(mensaje):
    assert not finanbot_ia_2.es_solicitud_documento(mensaje)


def test_reconoce_reporte_financiero_estandar_sin_interceptar_documento_personalizado():
    assert finanbot_ia_2.es_reporte_financiero_estandar('Hazme un reporte de este mes')
    assert not finanbot_ia_2.es_reporte_financiero_estandar(
        'Crea un informe para presentar mis gastos a mi arrendador',
    )


def test_chat_analiza_adjunto_y_guarda_el_enlace_en_el_historial(monkeypatch):
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    usuario = Usuario(nombre='Ana', correo='ana-chat-adjunto@example.com', contrasena_hash='hash')
    session.add(usuario)
    session.commit()
    adjunto_id = str(uuid.uuid4())
    conversacion_id = str(uuid.uuid4())
    session.add(ChatAdjunto(
        id=adjunto_id,
        usuario_id=usuario.id,
        nombre='balance octubre.pdf',
        ruta=f'chat/{usuario.id}/{adjunto_id}.pdf',
        mime_type='application/pdf',
        texto_extraido='Balance neto: $1.200.000',
    ))
    session.commit()
    recibido = {}

    class FinanBotFalso:
        _documento_generado = {
            'titulo': 'Presupuesto basado en el archivo',
            'formato': 'pdf',
            'contenido': 'Ingresos y gastos del archivo.',
        }

        def responder_con_acciones(self, mensaje, ctx, accion, conversacion_id=None):
            recibido['mensaje'] = mensaje
            return 'Resumen listo', []

    monkeypatch.setattr(chat_routes, 'sesiones', {})
    monkeypatch.setattr(chat_routes, 'FinanBotIA', FinanBotFalso)
    monkeypatch.setattr(chat_routes, 'preparar_adjuntos_chat', lambda *args: (
        'Contenido de balance octubre.pdf:\nBalance neto: $1.200.000', [],
    ))
    monkeypatch.setattr(chat_routes, '_cargar_contexto', lambda uid, db: {'usuario_id': uid})
    monkeypatch.setattr(chat_routes, '_actualizar_contexto', lambda *args: None)
    monkeypatch.setattr(chat_routes, '_agregar_disclaimers_si_necesario', lambda respuesta, mensaje: respuesta)
    monkeypatch.setattr(chat_routes, '_agregar_avisos_metas', lambda respuesta, ctx, accion: respuesta)

    try:
        resultado = chat_routes.mensaje(
            {'mensaje': 'Resume el archivo', 'adjuntos': [adjunto_id], 'conversacion_id': conversacion_id},
            str(usuario.id),
            session,
        )
        assert 'Balance neto: $1.200.000' in recibido['mensaje']
        assert resultado['respuesta'] == 'Resumen listo'
        assert resultado['documento'] is None
        mensaje_guardado = session.query(Chat).one().mensaje
        assert f'[[archivo:{adjunto_id}|balance%20octubre.pdf]]' in mensaje_guardado
    finally:
        session.close()
        engine.dispose()
